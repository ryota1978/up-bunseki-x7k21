"""
クレジット分析 シート

オリコCSV（全支払月）を勘定科目 × 支払月でクロス集計。
事業利用分のみ（私的（家族）を除外）を対象とする。
"""

from collections import defaultdict

from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

from ._common import (
    new_sheet, set_title, section, set_widths, write_header, write_row,
    note_block, month_cross_table, NORMAL_FONT, BOLD_FONT, SUBHEADER_FILL,
    BORDER, YEN, PCT,
)


SHEET_NAME = "クレジット分析"

EXCLUDE_FLAG = "私的（家族）"


def aggregate(cc):
    """勘定科目 × 支払月（事業利用分のみ）"""
    data = defaultdict(lambda: defaultdict(int))
    months = set()
    for r in cc:
        if r.get("biz_flag") == EXCLUDE_FLAG:
            continue
        if not r.get("amount"):
            continue
        acct = r.get("account") or "未分類"
        if acct.startswith("_"):
            continue
        m = r.get("payment_month")
        if not m:
            continue
        months.add(m)
        data[acct][m] += r["amount"]
    return {k: dict(v) for k, v in data.items()}, sorted(months)


def create_sheet(wb, records=None, cc=None):
    ws = new_sheet(wb, SHEET_NAME)
    set_title(ws, "💳 クレジットカード（オリコ）全期間分析", "A1:N1")

    cc = cc or []
    data, months = aggregate(cc)

    # ---------- 月別・科目別 ----------
    section(ws, 3, "■ 月別・科目別 内訳（事業利用分）")
    hdr, first, last, total_row, total_col = month_cross_table(
        ws, 4, "勘定科目", data, months=months
    )

    # ---------- 事業/私的 サマリ ----------
    sec2 = total_row + 3
    section(ws, sec2, "■ 事業 / 私的 の内訳")
    r = write_header(ws, sec2 + 1, ["区分", "件数", "金額", "シェア"])
    flags = defaultdict(lambda: [0, 0])
    for x in cc:
        if not x.get("amount"):
            continue
        f = x.get("biz_flag") or "不明"
        flags[f][0] += 1
        flags[f][1] += x["amount"]
    grand = sum(v[1] for v in flags.values())
    f_first = r
    for f, (n, amt) in sorted(flags.items(), key=lambda kv: -kv[1][1]):
        r = write_row(ws, r, [f, n, amt, (amt / grand if grand else 0)],
                      number_cols={3}, pct_cols={4})
    f_last = r - 1
    write_row(ws, r, ["合計", f"=SUM(B{f_first}:B{f_last})",
                      f"=SUM(C{f_first}:C{f_last})", 1],
              number_cols={3}, pct_cols={4}, bold=True, fill=SUBHEADER_FILL)

    # ---------- グラフ ----------
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "stacked"
    bar.overlap = 100
    bar.style = 11
    bar.title = "オリコ 科目別 月次推移（積上げ）"
    bar.height = 11
    bar.width = 22
    bar.y_axis.number_format = "#,##0"
    d = Reference(ws, min_col=2, min_row=hdr, max_col=1 + len(months), max_row=last)
    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    bar.add_data(d, titles_from_data=True, from_rows=True)
    bar.set_categories(cats)
    ws.add_chart(bar, f"A{r + 3}")

    pie = PieChart()
    pie.title = "オリコ 科目別シェア（全期間・事業分）"
    pie.height = 11
    pie.width = 14
    pd = Reference(ws, min_col=1 + len(months) + 1, min_row=first, max_row=last)
    pc = Reference(ws, min_col=1, min_row=first, max_row=last)
    pie.add_data(pd, titles_from_data=False)
    pie.set_categories(pc)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    ws.add_chart(pie, f"L{r + 3}")

    biz_total = sum(v for acc in data.values() for v in acc.values())
    note_block(ws, r + 26, [
        "📝 注釈",
        f"• 対象: オリコCSV {len(cc)}件のうち事業利用分（私的（家族）を除外）",
        f"• 事業分合計 ¥{biz_total:,}（{len(months)}ヶ月分の支払月データ）",
        "• 家族カードのガソリン・ETCは事業経費として計上しています",
        "• 家族カードの食事は「私的（家族）」として集計から除外しています",
        "• 2025/09 支払分のオリコCSVは未提供のため欠落しています",
    ])

    set_widths(ws, {"A": 18, **{get_column_letter(i): 13 for i in range(2, len(months) + 3)}})
    ws.freeze_panes = "B5"
    return ws
