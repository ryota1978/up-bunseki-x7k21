"""
Amazon購入分析 シート

amazon_products_v3.json（Dropbox上のAmazon領収書PDFから抽出した商品明細）を
カテゴリ別・配送先別に集計する。

⚠️ 重要な制約:
   Amazon領収書PDFの多くは「画像のみ（テキストレイヤーなし）」のため、
   商品名を抽出できたのは一部の領収書に限られる。
   金額はすべての領収書について判明しているので、
   「商品明細あり」「抽出不可」を明示的に分けて集計する。
"""

import re
from collections import defaultdict

from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

from config import TARGET_MONTHS
from parsers import categorize_amazon_item
from ._common import (
    new_sheet, set_title, section, set_widths, write_header, write_row,
    note_block, SUBHEADER_FILL, WARNING_FILL, YEN, PCT,
)


SHEET_NAME = "Amazon購入分析"

UNEXTRACTED = "その他(抽出不可・画像PDF)"

# parsers.categorize_amazon_item を補う追加パターン（parsers.py は変更しない）
EXTRA_CATEGORIES = [
    ("清掃・消耗品", [r"マジックリン", r"バスタブ", r"洗剤", r"キュキュット", r"スポンジ",
                     r"掃除機", r"クリーナー", r"タオル", r"ブラーバ", r"ラップ"]),
    ("電化製品・IT機器", [r"ケーブル", r"電源タップ", r"コンセント", r"充電", r"電池",
                         r"ＵＳＢ", r"USB", r"ハンガー", r"結束バンド", r"クリップ",
                         r"ＥＣＩ-", r"ECI-"]),
    ("医療・薬局備品", [r"第2類医薬品", r"第２類医薬品", r"絆創膏", r"ローション",
                       r"クリーム", r"シャンプー", r"歯磨き", r"シュミテクト",
                       r"医薬部外品"]),
    ("防災・安全用品", [r"アルミシート", r"サバイバル", r"防災", r"備蓄"]),
    ("飲食・休憩用品", [r"い・ろ・は・す", r"いろはす", r"お茶", r"ミネラルウォーター",
                       r"ティーバッグ"]),
    ("健康・運動用品", [r"ダンベル", r"筋膜ローラー", r"トレーニング"]),
    ("自転車・配達用品", [r"自転車", r"タイヤ", r"チューブ", r"タイヤレバー"]),
    ("オフィス家具", [r"ウォールポケット", r"コルクボード", r"ボード"]),
    ("文房具・事務用品", [r"消しゴム", r"マグネット", r"ポケット", r"ファイル", r"ラベル"]),
]


def categorize(name):
    """parsers の分類を先に使い、その他になったら追加パターンで再判定"""
    cat = categorize_amazon_item(name)
    if cat != "その他":
        return cat
    for c, pats in EXTRA_CATEGORIES:
        for p in pats:
            if re.search(p, name, re.IGNORECASE):
                return c
    return "その他"


def allocate(receipt):
    """
    領収書金額を商品行に按分する。
    価格が取れている行はその比率で、取れていなければ均等按分。
    """
    items = receipt["items"]
    total = receipt.get("receipt_amount") or 0
    if not items:
        return []
    weights = [(it.get("price") or 0) * (it.get("qty") or 1) for it in items]
    wsum = sum(weights)
    out = []
    for it, w in zip(items, weights):
        if wsum > 0:
            amt = round(total * w / wsum)
        else:
            amt = round(total / len(items))
        out.append((it, amt))
    return out


def aggregate(amazon, records):
    """カテゴリ別 / 配送先別 / 月別 / 配送先×カテゴリ を集計"""
    by_cat = defaultdict(lambda: [0, 0])       # [件数, 金額]
    by_ship = defaultdict(lambda: [0, 0])
    by_month = defaultdict(lambda: [0, 0])
    cross = defaultdict(lambda: defaultdict(int))
    rows = []
    examples = defaultdict(list)

    n_receipts = 0
    n_with_text = 0
    total_amount = 0

    for r in amazon or []:
        m = r.get("month")
        if m not in TARGET_MONTHS:
            continue
        n_receipts += 1
        amt = r.get("receipt_amount") or 0
        total_amount += amt
        ship = r.get("shipto") or "不明"
        by_ship[ship][0] += 1
        by_ship[ship][1] += amt
        by_month[m][0] += 1
        by_month[m][1] += amt

        if r.get("items"):
            n_with_text += 1
            for it, alloc in allocate(r):
                cat = categorize(it["name"])
                by_cat[cat][0] += 1
                by_cat[cat][1] += alloc
                cross[ship][cat] += alloc
                rows.append((m, r["filename"], ship, cat, it["name"],
                             it.get("qty") or 1, alloc))
                if len(examples[cat]) < 6:
                    examples[cat].append(it["name"][:60])
        else:
            by_cat[UNEXTRACTED][0] += 1
            by_cat[UNEXTRACTED][1] += amt
            cross[ship][UNEXTRACTED] += amt

    rows.sort(key=lambda x: (x[0], x[1]))
    return {
        "by_cat": dict(by_cat), "by_ship": dict(by_ship), "by_month": dict(by_month),
        "cross": {k: dict(v) for k, v in cross.items()}, "rows": rows,
        "examples": dict(examples), "n_receipts": n_receipts,
        "n_with_text": n_with_text, "total": total_amount,
    }


def create_sheet(wb, records=None, cc=None, amazon=None):
    ws = new_sheet(wb, SHEET_NAME)
    set_title(ws, "📦 Amazon購入分析（カテゴリ・配送先別）")

    if not amazon:
        ws["A3"] = "amazon_products_v3.json が見つからないため、商品明細の分析ができません。"
        ws["A3"].fill = WARNING_FILL
        ws["A4"] = "main_excel.py の load_data() が読み込むパスを確認してください。"
        return ws

    agg = aggregate(amazon, records)
    total = agg["total"]
    n_items = len(agg["rows"])

    # ---------- サマリ ----------
    section(ws, 3, "■ サマリ")
    r = 4
    coverage = agg["n_with_text"] / agg["n_receipts"] if agg["n_receipts"] else 0
    for label, val, fmt in [
        ("対象期間", f"{TARGET_MONTHS[0]} 〜 {TARGET_MONTHS[-1]}（{len(TARGET_MONTHS)}ヶ月）", None),
        ("領収書数", f"{agg['n_receipts']} 件", None),
        ("うち商品名を抽出できた領収書", f"{agg['n_with_text']} 件", None),
        ("商品名の抽出率", coverage, PCT),
        ("商品行 抽出数", f"{n_items} 件", None),
        ("総購入金額（領収書ベース）", total, YEN),
        ("月平均", total // len(TARGET_MONTHS), YEN),
    ]:
        ws.cell(row=r, column=1, value=label).fill = SUBHEADER_FILL
        c = ws.cell(row=r, column=2, value=val)
        if fmt:
            c.number_format = fmt
        r += 1
    ws.cell(row=r, column=1, value="金額の考え方").fill = SUBHEADER_FILL
    ws.cell(row=r, column=2,
            value="領収書総額を、商品行の単価×数量の比率で按分推定（個別単価は領収書PDF参照）")
    r += 1

    # ---------- カテゴリ別 ----------
    sec2 = r + 2
    section(ws, sec2, "■ カテゴリ別 推定金額")
    cr = write_header(ws, sec2 + 1, ["カテゴリ", "件数", "推定金額", "シェア"])
    c_first = cr
    for cat, (n, amt) in sorted(agg["by_cat"].items(), key=lambda kv: -kv[1][1]):
        cr = write_row(ws, cr, [cat, n, amt, (amt / total if total else 0)],
                       number_cols={3}, pct_cols={4},
                       fill=WARNING_FILL if cat == UNEXTRACTED else None)
    c_last = cr - 1
    write_row(ws, cr, ["合計", f"=SUM(B{c_first}:B{c_last})",
                       f"=SUM(C{c_first}:C{c_last})", 1],
              number_cols={3}, pct_cols={4}, bold=True, fill=SUBHEADER_FILL)

    # ---------- 配送先別 ----------
    sec3 = cr + 3
    section(ws, sec3, "■ 配送先別 金額")
    sr = write_header(ws, sec3 + 1, ["配送先", "領収書数", "金額", "シェア"])
    s_first = sr
    for ship, (n, amt) in sorted(agg["by_ship"].items(), key=lambda kv: -kv[1][1]):
        sr = write_row(ws, sr, [ship, n, amt, (amt / total if total else 0)],
                       number_cols={3}, pct_cols={4},
                       fill=WARNING_FILL if ship == "不明" else None)
    s_last = sr - 1
    write_row(ws, sr, ["合計", f"=SUM(B{s_first}:B{s_last})",
                       f"=SUM(C{s_first}:C{s_last})", 1],
              number_cols={3}, pct_cols={4}, bold=True, fill=SUBHEADER_FILL)

    # ---------- 月別 ----------
    sec4 = sr + 3
    section(ws, sec4, "■ 月別 購入金額")
    mr = write_header(ws, sec4 + 1, ["月度", "領収書数", "金額"])
    m_first = mr
    for m in TARGET_MONTHS:
        n, amt = agg["by_month"].get(m, (0, 0))
        mr = write_row(ws, mr, [m, n, amt], number_cols={3})
    write_row(ws, mr, ["合計", f"=SUM(B{m_first}:B{mr - 1})",
                       f"=SUM(C{m_first}:C{mr - 1})"],
              number_cols={3}, bold=True, fill=SUBHEADER_FILL)

    # ---------- 配送先 × カテゴリ ----------
    sec5 = mr + 3
    section(ws, sec5, "■ 配送先 × カテゴリ クロス集計（推定金額）")
    cats = [c for c, _ in sorted(agg["by_cat"].items(), key=lambda kv: -kv[1][1])]
    xhdr = sec5 + 1
    write_header(ws, xhdr, ["配送先"] + cats + ["合計"])
    xr = xhdr + 1
    x_first = xr
    from openpyxl.utils import get_column_letter
    ncol = 1 + len(cats)
    for ship in sorted(agg["cross"], key=lambda s: -agg["by_ship"][s][1]):
        vals = [ship] + [agg["cross"][ship].get(c, 0) for c in cats]
        vals.append(f"=SUM(B{xr}:{get_column_letter(ncol)}{xr})")
        xr = write_row(ws, xr, vals, number_cols=set(range(2, ncol + 2)))
    x_last = xr - 1
    write_row(ws, xr, ["合計"] + [
        f"=SUM({get_column_letter(j)}{x_first}:{get_column_letter(j)}{x_last})"
        for j in range(2, ncol + 2)
    ], number_cols=set(range(2, ncol + 2)), bold=True, fill=SUBHEADER_FILL)

    # ---------- グラフ ----------
    pie = PieChart()
    pie.title = "Amazon購入 カテゴリ別 金額シェア"
    pie.height = 11
    pie.width = 15
    pd = Reference(ws, min_col=3, min_row=c_first, max_row=c_last)
    pc = Reference(ws, min_col=1, min_row=c_first, max_row=c_last)
    pie.add_data(pd, titles_from_data=False)
    pie.set_categories(pc)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    ws.add_chart(pie, "F3")

    bar = BarChart()
    bar.type = "bar"
    bar.title = "配送先別 金額"
    bar.height = 10
    bar.width = 15
    bar.y_axis.number_format = "#,##0"
    bd = Reference(ws, min_col=3, min_row=s_first, max_row=s_last)
    bc = Reference(ws, min_col=1, min_row=s_first, max_row=s_last)
    bar.add_data(bd, titles_from_data=False)
    bar.set_categories(bc)
    ws.add_chart(bar, f"F{sec3}")

    # ---------- カテゴリ別 代表商品 ----------
    sec6 = xr + 3
    section(ws, sec6, "■ カテゴリ別 代表商品（実際に抽出できた購入品）")
    er = write_header(ws, sec6 + 1, ["カテゴリ", "代表的な購入商品（抜粋）"])
    for cat in cats:
        if cat == UNEXTRACTED:
            er = write_row(ws, er, [cat, "テキストレイヤーのない画像PDFのため商品名を抽出できなかった領収書"],
                           fill=WARNING_FILL)
            continue
        ex = agg["examples"].get(cat, [])
        er = write_row(ws, er, [cat, "／".join(ex) if ex else "-"])

    # ---------- 商品明細 ----------
    sec7 = er + 3
    section(ws, sec7, "■ 商品明細（抽出できた全行）")
    dr = write_header(ws, sec7 + 1,
                      ["月度", "ファイル名", "配送先", "カテゴリ", "商品名", "数量", "推定金額"])
    for row in agg["rows"]:
        dr = write_row(ws, dr, list(row), number_cols={7})

    # ---------- 注釈 ----------
    unext = agg["by_cat"].get(UNEXTRACTED, [0, 0])
    note_block(ws, dr + 2, [
        "📝 注釈・データの制約",
        f"• 領収書 {agg['n_receipts']}件 のうち、商品名を抽出できたのは {agg['n_with_text']}件（{coverage:.0%}）です",
        f"• 残り {unext[0]}件（¥{unext[1]:,}）は PDF が画像のみでテキストレイヤーを持たないため、",
        "  商品名が取得できませんでした。金額・月度・配送先の一部は判明しています。",
        "• これらから商品名を得るには OCR 処理が必要です（本システムでは未実施）",
        "• ファイル名が小文字 amazon の領収書はブラウザ印刷PDFでテキストあり、",
        "  大文字 Amazon のものはスクリーンショット由来で画像のみ、という傾向があります",
        "• 金額は領収書総額を商品行に按分した推定値です。厳密な単価は領収書PDFを参照してください",
        "• 配送先「不明」は、テキストが取れず配送先も判定できなかった領収書です",
    ])

    set_widths(ws, {"A": 26, "B": 34, "C": 24, "D": 26, "E": 70, "F": 10, "G": 14,
                    "H": 14, "I": 14, "J": 14, "K": 14, "L": 14, "M": 14})
    return ws
