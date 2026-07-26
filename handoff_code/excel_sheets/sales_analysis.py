"""
売上分析 シート

取引先（決済サービス・委託先・店舗）別 × 月別 の売上推移。
売上系の勘定科目（売上_調剤 / 売上_業務委託 / 売上集計）を対象とする。
"""

import re
from collections import defaultdict

from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from config import TARGET_MONTHS
from ._common import (
    new_sheet, set_title, section, set_widths, month_cross_table, note_block,
)


SHEET_NAME = "売上分析"

# ファイル名 → 売上取引先名 の判定パターン（上から順に評価）
SALES_VENDOR_PATTERNS = [
    ("PayPay", r"paypay|ペイペイ"),
    ("Airペイ クレジット", r"エアペイ(?!QR)|Airペイ(?!QR)|エアペイ取引集計"),
    ("Airペイ QR", r"エアペイQR|AirペイQR|エアペイ.*QR"),
    ("Airメイト", r"エアメイト|Airメイト"),
    ("メルペイ", r"メル[ぺペ]イ"),
    ("メルカリ", r"メルカリ"),
    ("名城大学", r"名城大学"),
    ("鈴鹿医療科学大学", r"鈴鹿医療科学大学"),
    ("度会特別支援学校", r"度会特別支援学校"),
    ("タイガー薬局", r"タイガー薬局"),
    ("リバイバルドラッグ", r"リバイバル|リバドラ"),
    ("ユニスマイル", r"ユニスマイル"),
    ("射和ロッカー", r"射和ロッカー"),
    ("松阪市役所", r"松阪市役所"),
    ("ひかり調剤薬局", r"ひかり調剤薬局"),
    ("ひかり薬局(松阪/射和)", r"ひかり薬局|射和店"),
    ("ひかりファーマシー", r"ひかりファーマシー"),
    ("ひかりハート薬局", r"ひかりハート"),
]


def classify_sales_vendor(filename):
    """売上レコードのファイル名から取引先名を判定"""
    fn = filename or ""
    for name, pat in SALES_VENDOR_PATTERNS:
        if re.search(pat, fn, re.IGNORECASE):
            return name
    return "その他売上"


def aggregate(records):
    """売上系レコードを 取引先 × 月 に集計"""
    data = defaultdict(lambda: defaultdict(int))
    for r in records:
        acct = r.get("account") or ""
        if not acct.startswith("売上"):
            continue
        m = r.get("month")
        if m not in TARGET_MONTHS or not r.get("amount"):
            continue
        data[classify_sales_vendor(r.get("filename"))][m] += r["amount"]
    return {k: dict(v) for k, v in data.items()}


def create_sheet(wb, records, cc=None):
    ws = new_sheet(wb, SHEET_NAME)
    set_title(ws, "💰 売上分析（全期間）")

    data = aggregate(records)
    section(ws, 3, "■ 取引先別 × 月別 売上推移")

    hdr, first, last, total_row, total_col = month_cross_table(
        ws, 4, "取引先", data
    )

    grand = sum(sum(v.values()) for v in data.values())

    # 積上げ棒グラフ（参考ブックはグラフ0だが、月次構成が一目で分かるので追加）
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.style = 12
    chart.title = "取引先別 月次売上（積上げ）"
    chart.height = 12
    chart.width = 26
    chart.y_axis.number_format = "#,##0"
    data_ref = Reference(ws, min_col=2, min_row=hdr,
                         max_col=1 + len(TARGET_MONTHS), max_row=last)
    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    chart.add_data(data_ref, titles_from_data=True, from_rows=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{total_row + 3}")

    note_row = total_row + 30
    note_block(ws, note_row, [
        "📝 注釈",
        f"• 売上系勘定科目（売上_調剤 / 売上_業務委託 / 売上集計）の合計 ¥{grand:,}",
        "• 取引先名は領収書・集計PDFのファイル名から判定しています",
        "• PayPay は店舗別の内訳が取れるため「お客様決済_店舗別売上」シートに詳細があります",
        "• Airペイ系は全店舗合算・店舗別が混在しており、店舗単位の分解はできていません",
    ])

    set_widths(ws, {"A": 26, **{get_column_letter(2 + i): 13 for i in range(len(TARGET_MONTHS) + 1)}})
    ws.freeze_panes = "B5"
    return ws
