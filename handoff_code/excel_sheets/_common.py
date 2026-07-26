"""
シート生成モジュール共通ヘルパー

config.STYLE を単一の情報源として、各シートモジュールが使う
フォント / 塗り / 罫線 / 書式 を提供する。
色やフォント名をここ以外でハードコードしないこと。
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import STYLE, TARGET_MONTHS


# ============ 共通スタイル定数 ============
HEADER_FILL = PatternFill("solid", start_color=STYLE["primary_color"])
HEADER_FONT = Font(name=STYLE["font_normal"], bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=STYLE["font_normal"], bold=True, size=14, color=STYLE["primary_color"])
SECTION_FONT = Font(name=STYLE["font_normal"], bold=True, size=12, color=STYLE["primary_color"])
NORMAL_FONT = Font(name=STYLE["font_normal"], size=10)
BOLD_FONT = Font(name=STYLE["font_normal"], size=10, bold=True)
NOTE_FONT = Font(name=STYLE["font_normal"], size=9)
ACCENT_FONT = Font(name=STYLE["font_normal"], size=12, bold=True, color=STYLE["accent_color"])
SUCCESS_FONT = Font(name=STYLE["font_normal"], size=12, bold=True, color=STYLE["success_color"])

SUBHEADER_FILL = PatternFill("solid", start_color=STYLE["sub_header"])
WARNING_FILL = PatternFill("solid", start_color=STYLE["warning_color"])
URIAGE_FILL = PatternFill("solid", start_color=STYLE["urihi_fill"])

BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

YEN = "#,##0;-#,##0;-"
YEN_MARK = "¥#,##0"
PCT = "0.0%"


def new_sheet(wb, name):
    """同名シートがあれば削除してから作成（再実行時の重複防止）"""
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def set_title(ws, text, span="A1:L1"):
    """シート左上のタイトル行を設定"""
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    try:
        ws.merge_cells(span)
    except Exception:
        pass


def section(ws, row, text):
    """■ セクション見出しを書き込む"""
    c = ws.cell(row=row, column=1, value=text)
    c.font = SECTION_FONT
    return c


def write_header(ws, row, headers, start_col=1):
    """ヘッダー行（濃紺塗り・白太字）を書き込む"""
    for i, h in enumerate(headers, start_col):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    return row + 1


def write_row(ws, row, values, start_col=1, number_cols=(), fill=None,
              bold=False, pct_cols=(), fmt=YEN):
    """
    1行分のデータを書き込む。

    Args:
        number_cols: 数値書式を当てる列番号（絶対列番号）の集合
        pct_cols:    パーセント書式を当てる列番号の集合
    """
    font = BOLD_FONT if bold else NORMAL_FONT
    for i, v in enumerate(values, start_col):
        c = ws.cell(row=row, column=i, value=v)
        c.font = font
        c.border = BORDER
        if fill is not None:
            c.fill = fill
        if i in pct_cols:
            c.number_format = PCT
        elif i in number_cols:
            c.number_format = fmt
    return row + 1


def note_block(ws, start_row, lines, col=1):
    """📝 注釈などのフリーテキストブロックを書き込む"""
    r = start_row
    for line in lines:
        c = ws.cell(row=r, column=col, value=line)
        c.font = NOTE_FONT if line.startswith(("  ", "•", "・", "□", "→")) else NORMAL_FONT
        r += 1
    return r


def set_widths(ws, widths):
    """{"A": 12, ...} または {1: 12, ...} で列幅を設定"""
    for k, w in widths.items():
        col = k if isinstance(k, str) else get_column_letter(k)
        ws.column_dimensions[col].width = w


def month_cross_table(ws, start_row, row_label, data, months=None,
                      total_label="合計", sort_desc=True):
    """
    「行ラベル × 月 + 合計」のクロス集計表を、合計列/合計行に
    実際のExcel数式（=SUM(...)）を入れた形で書き出す共通処理。

    Args:
        start_row:  ヘッダー行の行番号
        row_label:  A列ヘッダーの見出し（例 "取引先"）
        data:       {row_name: {"YYYY/MM": amount}}
        months:     月リスト（省略時 config.TARGET_MONTHS）

    Returns:
        (header_row, first_data_row, last_data_row, total_row, total_col_letter)
    """
    months = months or TARGET_MONTHS
    hdr = start_row
    write_header(ws, hdr, [row_label] + list(months) + [total_label])

    n_cols = 1 + len(months) + 1
    total_col = get_column_letter(n_cols)

    items = list(data.items())
    if sort_desc:
        items.sort(key=lambda kv: -sum(kv[1].values()))

    r = hdr + 1
    first = r
    for name, monthly in items:
        ws.cell(row=r, column=1, value=name).font = NORMAL_FONT
        ws.cell(row=r, column=1).border = BORDER
        for j, m in enumerate(months, 2):
            c = ws.cell(row=r, column=j, value=monthly.get(m, 0))
            c.number_format = YEN
            c.font = NORMAL_FONT
            c.border = BORDER
        # 合計列は実数式
        c = ws.cell(row=r, column=n_cols,
                    value=f"=SUM(B{r}:{get_column_letter(n_cols - 1)}{r})")
        c.number_format = YEN
        c.font = BOLD_FONT
        c.border = BORDER
        r += 1
    last = r - 1

    # 合計行も実数式
    total_row = r
    tc = ws.cell(row=total_row, column=1, value=total_label)
    tc.font = BOLD_FONT
    tc.fill = SUBHEADER_FILL
    tc.border = BORDER
    for j in range(2, n_cols + 1):
        col = get_column_letter(j)
        c = ws.cell(row=total_row, column=j, value=f"=SUM({col}{first}:{col}{last})")
        c.number_format = YEN
        c.font = BOLD_FONT
        c.fill = SUBHEADER_FILL
        c.border = BORDER

    return hdr, first, last, total_row, total_col
