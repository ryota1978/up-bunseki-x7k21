"""
消耗品_内訳分析 シート

消耗品費を 取引先 × 月 で集計（領収書＋クレジット合算）。
あわせて水道光熱費（グリムス）の月次推移も併記する。
"""

import re
from collections import defaultdict

from openpyxl.utils import get_column_letter

from config import TARGET_MONTHS
from ._common import (
    new_sheet, set_title, section, set_widths, write_header, write_row,
    note_block, month_cross_table, SUBHEADER_FILL, YEN,
)


SHEET_NAME = "消耗品_内訳分析"

# 取引先の正規化パターン
VENDOR_ALIASES = [
    ("Amazon", [r"amazon", r"amzon", r"amaon", r"ＡＭＡＺＯＮ"]),
    ("アスクル(ソロエル)", [r"ソロエル", r"アスクル", r"アリーナ"]),
    ("ヨドバシ", [r"ヨドバシ"]),
    ("ビックカメラ", [r"ビックカメラ", r"ビツクカメラ"]),
    ("ヤマダデンキ", [r"ヤマダ"]),
    ("ケーズデンキ", [r"ケーズ", r"ケ−ズ"]),
    ("楽天", [r"楽天", r"Rakuten", r"ラクテン"]),
    ("メルカリ", [r"メルカリ"]),
    ("ウエルシア", [r"ウエルシア"]),
    ("ミライヤ書店", [r"ミライヤ"]),
    ("ダイソー", [r"ダイソー", r"ダイソ−"]),
    ("オフィスコム", [r"オフィスコム"]),
    ("日本郵便", [r"日本郵便", r"ニツポンユウビン", r"郵便"]),
    ("高島屋", [r"高島屋", r"タカシマヤ"]),
    ("カメラのキタムラ", [r"キタムラ"]),
    ("セカンドストリート", [r"セカンドストリート"]),
]


def normalize(*fields):
    blob = " ".join(f for f in fields if f)
    for name, pats in VENDOR_ALIASES:
        for p in pats:
            if re.search(p, blob, re.IGNORECASE):
                return name
    return "その他"


def aggregate(records, cc):
    data = defaultdict(lambda: defaultdict(int))
    counts = defaultdict(int)
    for r in records or []:
        if r.get("account") != "消耗品費" or not r.get("amount"):
            continue
        m = r.get("month")
        if m not in TARGET_MONTHS:
            continue
        v = normalize(r.get("tag"), r.get("vendor"), r.get("filename"))
        data[v][m] += r["amount"]
        counts[v] += 1
    for r in cc or []:
        if r.get("account") != "消耗品費" or not r.get("amount"):
            continue
        if r.get("biz_flag") == "私的（家族）":
            continue
        m = r.get("payment_month")
        if m not in TARGET_MONTHS:
            continue
        v = normalize(r.get("tag"), r.get("vendor"))
        # クレジット側の「その他」は出所が分かるよう区別
        data[v if v != "その他" else "その他(CC)"][m] += r["amount"]
        counts[v if v != "その他" else "その他(CC)"] += 1
    return {k: dict(v) for k, v in data.items()}, dict(counts)


def grimms_monthly(cc):
    """水道光熱費（グリムス）の支払月別 件数/金額"""
    out = defaultdict(lambda: [0, 0])
    for r in cc or []:
        if r.get("account") != "水道光熱費" or not r.get("amount"):
            continue
        if r.get("biz_flag") == "私的（家族）":
            continue
        m = r.get("payment_month") or "-"
        out[m][0] += 1
        out[m][1] += r["amount"]
    return dict(out)


def create_sheet(wb, records=None, cc=None):
    ws = new_sheet(wb, SHEET_NAME)
    set_title(ws, "🛒 消耗品費 詳細内訳")

    data, counts = aggregate(records, cc)

    section(ws, 3, "■ 取引先別 月次推移（領収書＋クレジット 合算）")
    hdr, first, last, total_row, total_col = month_cross_table(
        ws, 4, "取引先", data
    )

    # ---------- 水道光熱費 ----------
    sec2 = total_row + 3
    section(ws, sec2, "■ 水道光熱費（電力会社グリムス）月次推移")
    gm = grimms_monthly(cc)
    gr = write_header(ws, sec2 + 1, ["支払月", "件数", "合計"])
    g_first = gr
    for m in sorted(gm):
        n, amt = gm[m]
        gr = write_row(ws, gr, [m, n, amt], number_cols={3})
    write_row(ws, gr, ["合計", f"=SUM(B{g_first}:B{gr - 1})",
                       f"=SUM(C{g_first}:C{gr - 1})"],
              number_cols={3}, bold=True, fill=SUBHEADER_FILL)

    total = sum(sum(v.values()) for v in data.values())
    amazon = sum(data.get("Amazon", {}).values())
    note_block(ws, gr + 3, [
        "📝 注釈",
        f"• 消耗品費 合計 ¥{total:,}",
        f"• うち Amazon ¥{amazon:,}（{(amazon / total if total else 0):.1%}）"
        " → 商品明細は「Amazon購入分析」シート参照",
        "• 取引先名は領収書のファイル名／クレジット取引先名から正規化しています",
        "• 「その他(CC)」はクレジット側で個別に名寄せできなかった小売店です",
        "• 電気代の店舗別内訳は「電気代_店舗別分析」シートを参照してください",
    ])

    set_widths(ws, {"A": 22, **{get_column_letter(i): 13 for i in range(2, len(TARGET_MONTHS) + 3)}})
    ws.freeze_panes = "B5"
    return ws
