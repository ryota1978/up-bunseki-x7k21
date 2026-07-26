"""
保険料_内訳分析 シート

日薬保険は amount_corrections.NICHIYAKU_BREAKDOWN の4店舗内訳をそのまま使用。
三井住友海上はクレジット明細（account=保険料）から実額を取得する。
"""

from collections import defaultdict

from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList

from config import STORES
from amount_corrections import NICHIYAKU_BREAKDOWN
from ._common import (
    new_sheet, set_title, section, set_widths, write_header, write_row,
    note_block, SUBHEADER_FILL, WARNING_FILL, YEN, PCT,
)


SHEET_NAME = "保険料_内訳分析"


def collect_msa(cc):
    """三井住友海上のクレジット明細を抽出"""
    rows = []
    for r in cc or []:
        if r.get("account") != "保険料" or not r.get("amount"):
            continue
        if "三井住友" not in (r.get("vendor") or ""):
            continue
        rows.append(r)
    rows.sort(key=lambda x: -x["amount"])
    return rows


def create_sheet(wb, records=None, cc=None):
    ws = new_sheet(wb, SHEET_NAME)
    set_title(ws, "🛡️ 保険料 詳細内訳", "A1:G1")

    msa_rows = collect_msa(cc)
    msa_total = sum(r["amount"] for r in msa_rows)
    nichiyaku_total = sum(x[4] for x in NICHIYAKU_BREAKDOWN)
    grand = msa_total + nichiyaku_total
    n_stores = len(STORES)

    # ---------- 全体サマリ ----------
    section(ws, 3, "■ 保険料 全体サマリ")
    r = write_header(ws, 4, ["保険種別", "保険会社", "支払日", "金額", "支払方式", "備考"])
    sum_first = r
    labels = ["業務用火災保険（推定）", "追加保険（推定）"]
    for i, row in enumerate(msa_rows):
        r = write_row(ws, r, [
            labels[i] if i < len(labels) else f"三井住友海上 その他({i + 1})",
            "三井住友海上", row.get("date", ""), row["amount"], "年払い一括",
            "4店舗物件保険・大口契約" if i == 0 else "賠償責任/休業補償等",
        ], number_cols={4})
    r = write_row(ws, r, [
        "薬剤師会員保険", f"日薬国保（{n_stores}店舗分）", "-", nichiyaku_total,
        "都度払い", f"{n_stores}店舗合算（下に内訳）",
    ], number_cols={4})
    sum_last = r - 1
    write_row(ws, r, ["合計", "", "", f"=SUM(D{sum_first}:D{sum_last})", "", ""],
              number_cols={4}, bold=True, fill=SUBHEADER_FILL)
    pie_first, pie_last = sum_first, sum_last

    # ---------- 日薬保険 内訳 ----------
    sec2 = r + 3
    section(ws, sec2, f"■ 日薬保険 {n_stores}店舗別内訳（領収書 ¥{nichiyaku_total:,}）")
    nr = write_header(ws, sec2 + 1, ["管理番号", "薬局", "加入者名", "区分", "保険料", "課税"])
    n_first = nr
    for mgmt_id, store, person, cls, amount in NICHIYAKU_BREAKDOWN:
        taxed = "課税" if cls == "10%課税" else "非課税"
        label = f"薬局{cls}" if cls in ("E", "F") else cls
        nr = write_row(ws, nr, [mgmt_id, store, person, label, amount, taxed],
                       number_cols={5})
    n_last = nr - 1
    write_row(ws, nr, ["", "合計", "", "", f"=SUM(E{n_first}:E{n_last})", ""],
              number_cols={5}, bold=True, fill=SUBHEADER_FILL)

    # ---------- 三井住友海上 内訳 ----------
    sec3 = nr + 3
    section(ws, sec3, f"■ 三井住友海上 内訳（¥{msa_total:,}）")
    ws.cell(row=sec3 + 1, column=1,
            value="⚠️ クレジット明細では数件にまとまっているのみで、詳細内訳は保険証券の確認が必要です"
            ).fill = WARNING_FILL
    mr = write_header(ws, sec3 + 3,
                      ["支払日", "金額", "シェア", "推定内容", "月割（年÷12）"])
    m_first = mr
    guesses = ["業務用火災保険（4店舗合算）と推定",
               "賠償責任・休業補償・什器備品等の追加保険"]
    for i, row in enumerate(msa_rows):
        mr = write_row(ws, mr, [
            row.get("date", ""), row["amount"],
            (row["amount"] / msa_total if msa_total else 0),
            guesses[i] if i < len(guesses) else "内訳不明（保険証券要確認）",
            row["amount"] // 12,
        ], number_cols={2, 5}, pct_cols={3})
    m_last = mr - 1
    write_row(ws, mr, ["合計", f"=SUM(B{m_first}:B{m_last})", 1, "",
                       f"=SUM(E{m_first}:E{m_last})"],
              number_cols={2, 5}, pct_cols={3}, bold=True, fill=SUBHEADER_FILL)

    # ---------- 1店舗あたり ----------
    sec4 = mr + 3
    section(ws, sec4, f"■ 1店舗あたりの保険料 月割推定（{n_stores}店舗想定）")
    br = write_header(ws, sec4 + 1,
                      [f"保険種類", f"{n_stores}店舗合計（年）", "1店舗（年）", "1店舗（月）"])
    b_first = br
    for i, row in enumerate(msa_rows):
        label = "三井住友海上（火災保険）" if i == 0 else "三井住友海上（追加保険）"
        br = write_row(ws, br, [label, row["amount"], row["amount"] // n_stores,
                                row["amount"] // n_stores // 12], number_cols={2, 3, 4})
    br = write_row(ws, br, ["日薬保険", nichiyaku_total, nichiyaku_total // n_stores,
                            nichiyaku_total // n_stores // 12], number_cols={2, 3, 4})
    b_last = br - 1
    write_row(ws, br, ["合計", f"=SUM(B{b_first}:B{b_last})",
                       f"=SUM(C{b_first}:C{b_last})", f"=SUM(D{b_first}:D{b_last})"],
              number_cols={2, 3, 4}, bold=True, fill=SUBHEADER_FILL)

    # ---------- グラフ ----------
    pie = PieChart()
    pie.title = "保険料 内訳シェア"
    pie.height = 11
    pie.width = 14
    pd = Reference(ws, min_col=4, min_row=pie_first, max_row=pie_last)
    pc = Reference(ws, min_col=1, min_row=pie_first, max_row=pie_last)
    pie.add_data(pd, titles_from_data=False)
    pie.set_categories(pc)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    ws.add_chart(pie, "H4")

    # ---------- 所感 ----------
    per_store_month = grand // n_stores // 12
    note_block(ws, br + 3, [
        "📝 保険料についての所感",
        f"• 保険料 合計 ¥{grand:,}（三井住友海上 ¥{msa_total:,} ＋ 日薬保険 ¥{nichiyaku_total:,}）",
        f"• 1店舗あたり月 約¥{per_store_month:,} 程度",
        "• 日薬保険の区分Fは ひかり調剤薬局のみ（補償拡大版と推定）、他3店舗は区分E",
        "• 三井住友海上の内訳詳細は保険証券で確認を推奨",
        "• 保険証券が入手できれば、店舗別の保険金額・補償内容が明確になります",
        "• 年一括払いは資金繰り負担が大きいため、月払に変更可能か保険代理店に相談余地あり",
    ])

    set_widths(ws, {"A": 30, "B": 30, "C": 16, "D": 34, "E": 16, "F": 12, "G": 12})
    return ws
